import os
import io
import re
import json
import html
import queue
import base64
import asyncio
import logging
import threading
from contextlib import contextmanager
from datetime import datetime, date, time, timedelta
from zoneinfo import ZoneInfo

import requests
import psycopg2
from psycopg2.pool import ThreadedConnectionPool
from requests.auth import HTTPDigestAuth
from flask import Flask, request
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from telegram import (Update, InlineKeyboardButton, InlineKeyboardMarkup,
                      ReplyKeyboardMarkup)
from telegram.error import BadRequest
from telegram.ext import (Application, CommandHandler, CallbackQueryHandler,
                          MessageHandler, ChatMemberHandler, filters, ContextTypes)

load_dotenv()
logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"), format="%(asctime)s %(levelname)s %(message)s")
for _n in ("httpx", "httpcore", "urllib3", "werkzeug", "waitress", "apscheduler"):
    logging.getLogger(_n).setLevel(logging.WARNING)
log = logging.getLogger("davomat")

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

# ===================== SOZLAMALAR =====================
TG_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
# Bir nechta chat bo'lsa vergul bilan: 12345,-100987
TG_CHAT_IDS = [x.strip() for x in os.getenv("TELEGRAM_CHAT_ID", "").split(",") if x.strip()]
HIK_IP = os.getenv("HIKVISION_IP")
HIK_USER = os.getenv("HIKVISION_USER")
HIK_PASS = os.getenv("HIKVISION_PASSWORD")
PORT = int(os.getenv("PORT", "8080"))

TZ = ZoneInfo(os.getenv("TZ_NAME", "Asia/Tashkent"))
GRACE_MINUTES = int(os.getenv("GRACE_MINUTES", "0"))          # kechikish uchun imtiyoz
RETENTION_DAYS = int(os.getenv("RETENTION_DAYS", "180"))      # eski yozuvlar o'chiriladi
WEEKEND_DAYS = {int(x) for x in os.getenv("WEEKEND_DAYS", "6").split(",") if x.strip()}  # 6 = yakshanba
SUPER_ADMINS = {int(x) for x in os.getenv("ADMIN_IDS", "").replace(" ", "").split(",") if x}


def parse_hhmm(s: str):
    s = s.strip().replace(".", ":")
    m = re.fullmatch(r"(\d{1,2})(?::(\d{1,2}))?", s)
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2) or 0)
    if h > 23 or mi > 59:
        return None
    return time(h, mi)


DEFAULT_START = parse_hhmm(os.getenv("DEFAULT_START", "09:00")) or time(9, 0)

WEEKDAYS = ["Dushanba", "Seshanba", "Chorshanba", "Payshanba", "Juma", "Shanba", "Yakshanba"]


def now_local() -> datetime:
    return datetime.now(TZ).replace(tzinfo=None)


def esc(s) -> str:
    return html.escape(str(s))


def fmt_min(m) -> str:
    m = int(m or 0)
    return f"{m} daq" if m < 60 else f"{m // 60} soat {m % 60} daq"


# ===================== BAZA (connection pool) =====================
_pool = None
_pool_lock = threading.Lock()


def get_pool():
    global _pool
    with _pool_lock:
        if _pool is None:
            _pool = ThreadedConnectionPool(
                1, 3,
                host=os.getenv("DB_HOST"), port=os.getenv("DB_PORT", "5432"), database=os.getenv("DB_NAME"),
                user=os.getenv("DB_USER"), password=os.getenv("DB_PASS"),
            )
        return _pool


@contextmanager
def db():
    p = get_pool()
    conn = p.getconn()
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        p.putconn(conn)


def init_db():
    with db() as c, c.cursor() as cur:
        cur.execute("""
            CREATE TABLE IF NOT EXISTS attendance (
                id SERIAL PRIMARY KEY,
                employee_name TEXT,
                event_type TEXT,
                event_time TIMESTAMP,
                late_minutes INT DEFAULT 0
            )""")
        cur.execute("ALTER TABLE attendance ADD COLUMN IF NOT EXISTS emp_no TEXT")
        cur.execute("ALTER TABLE attendance ADD COLUMN IF NOT EXISTS serial_no BIGINT")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_att_time ON attendance(event_time)")
        cur.execute("CREATE INDEX IF NOT EXISTS idx_att_emp_time ON attendance(emp_no, event_time)")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS employees (
                emp_no TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                work_start TIME NOT NULL,
                active BOOLEAN NOT NULL DEFAULT TRUE
            )""")
        cur.execute("ALTER TABLE employees ADD COLUMN IF NOT EXISTS department TEXT")
        cur.execute("ALTER TABLE employees ADD COLUMN IF NOT EXISTS job_title TEXT")
        cur.execute("ALTER TABLE employees ADD COLUMN IF NOT EXISTS updated_at TIMESTAMPTZ NOT NULL DEFAULT now()")
        cur.execute("CREATE TABLE IF NOT EXISTS admins (tg_id BIGINT PRIMARY KEY, name TEXT)")
        cur.execute("ALTER TABLE admins ADD COLUMN IF NOT EXISTS role TEXT NOT NULL DEFAULT 'admin'")
        cur.execute("CREATE TABLE IF NOT EXISTS chats (chat_id BIGINT PRIMARY KEY, title TEXT)")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS config_version (
                id INT PRIMARY KEY DEFAULT 1 CHECK (id=1),
                version BIGINT NOT NULL DEFAULT 1
            )""")
        cur.execute("INSERT INTO config_version (id, version) VALUES (1, 1) ON CONFLICT DO NOTHING")


# ===================== KESH (DB ga ortiqcha so'rov bormasligi uchun) =====================
EMP = {}        # emp_no -> {"name", "start", "active"}
ADMINS = set()
CHATS = set()
_last_config_version = 0
LOCK = threading.Lock()



def bump_config_version(cur):
    cur.execute("""
        INSERT INTO config_version (id, version) VALUES (1, 1)
        ON CONFLICT (id) DO UPDATE SET version = config_version.version + 1
    """)


def load_caches():
    global _last_config_version
    with db() as c, c.cursor() as cur:
        cur.execute("SELECT emp_no, name, work_start, active FROM employees")
        emps = {r[0]: {"name": r[1], "start": r[2], "active": r[3]} for r in cur.fetchall()}
        cur.execute("SELECT tg_id FROM admins")
        admins = {r[0] for r in cur.fetchall()}
        cur.execute("SELECT chat_id FROM chats")
        chats = {r[0] for r in cur.fetchall()}
        cur.execute("SELECT version FROM config_version WHERE id=1")
        ver_row = cur.fetchone()
        if ver_row:
            _last_config_version = ver_row[0]
    with LOCK:
        EMP.clear(); EMP.update(emps)
        ADMINS.clear(); ADMINS.update(admins)
        CHATS.clear(); CHATS.update(chats)


def check_config_version():
    """Web yoki boshqa joydan o'zgargan xodimlar/adminlarni 30 soniyada yangilash"""
    global _last_config_version
    try:
        with db() as c, c.cursor() as cur:
            cur.execute("SELECT version FROM config_version WHERE id=1")
            ver_row = cur.fetchone()
            if ver_row and ver_row[0] != _last_config_version:
                log.info("Config versiyasi o'zgardi (%s -> %s), keshlar yangilanmoqda...",
                         _last_config_version, ver_row[0])
                _last_config_version = ver_row[0]
                cur.execute("SELECT emp_no, name, work_start, active FROM employees")
                emps = {r[0]: {"name": r[1], "start": r[2], "active": r[3]} for r in cur.fetchall()}
                cur.execute("SELECT tg_id FROM admins")
                admins = {r[0] for r in cur.fetchall()}
                cur.execute("SELECT chat_id FROM chats")
                chats = {r[0] for r in cur.fetchall()}
                with LOCK:
                    EMP.clear(); EMP.update(emps)
                    ADMINS.clear(); ADMINS.update(admins)
                    CHATS.clear(); CHATS.update(chats)
    except Exception as e:
        log.warning("Config versiya tekshirishda xatolik: %s", e)


def is_admin(uid: int) -> bool:
    return uid in SUPER_ADMINS or uid in ADMINS


# ---- xodimlar ----
def upsert_employee(emp_no, name, start: time):
    with db() as c, c.cursor() as cur:
        cur.execute("""
            INSERT INTO employees (emp_no, name, work_start, active) VALUES (%s,%s,%s,TRUE)
            ON CONFLICT (emp_no) DO UPDATE
            SET name = EXCLUDED.name, work_start = EXCLUDED.work_start, active = TRUE
        """, (emp_no, name, start))
        bump_config_version(cur)
    with LOCK:
        EMP[emp_no] = {"name": name, "start": start, "active": True}


def set_work_start(emp_no, start: time):
    with db() as c, c.cursor() as cur:
        cur.execute("UPDATE employees SET work_start=%s WHERE emp_no=%s", (start, emp_no))
        bump_config_version(cur)
    with LOCK:
        if emp_no in EMP:
            EMP[emp_no]["start"] = start


def set_active(emp_no, active: bool):
    with db() as c, c.cursor() as cur:
        cur.execute("UPDATE employees SET active=%s WHERE emp_no=%s", (active, emp_no))
        bump_config_version(cur)
    with LOCK:
        if emp_no in EMP:
            EMP[emp_no]["active"] = active


# ---- adminlar / chatlar ----
def add_admin(tg_id: int, name: str):
    with db() as c, c.cursor() as cur:
        cur.execute("""INSERT INTO admins (tg_id, name) VALUES (%s,%s)
                       ON CONFLICT (tg_id) DO UPDATE SET name=EXCLUDED.name""", (tg_id, name))
        bump_config_version(cur)
    with LOCK:
        ADMINS.add(tg_id)


def remove_admin(tg_id: int):
    with db() as c, c.cursor() as cur:
        cur.execute("DELETE FROM admins WHERE tg_id=%s", (tg_id,))
        bump_config_version(cur)
    with LOCK:
        ADMINS.discard(tg_id)


def list_admins():
    with db() as c, c.cursor() as cur:
        cur.execute("SELECT tg_id, COALESCE(name,'') FROM admins ORDER BY tg_id")
        return cur.fetchall()


def add_chat(chat_id, title):
    with db() as c, c.cursor() as cur:
        cur.execute("""INSERT INTO chats (chat_id, title) VALUES (%s,%s)
                       ON CONFLICT (chat_id) DO UPDATE SET title=EXCLUDED.title""", (chat_id, title))
        bump_config_version(cur)
    with LOCK:
        CHATS.add(chat_id)


def remove_chat(chat_id):
    with db() as c, c.cursor() as cur:
        cur.execute("DELETE FROM chats WHERE chat_id=%s", (chat_id,))
        bump_config_version(cur)
    with LOCK:
        CHATS.discard(chat_id)


def cleanup_old():
    cutoff = now_local() - timedelta(days=RETENTION_DAYS)
    with db() as c, c.cursor() as cur:
        # daily_attendance jadvali bo'lsa, yig'ilgan kunlargacha bo'lgan sanani tekshiramiz
        try:
            cur.execute("SELECT to_regclass('daily_attendance')")
            reg = cur.fetchone()
            if reg and reg[0]:
                cur.execute("SELECT MAX(day) FROM daily_attendance")
                max_agg = cur.fetchone()
                if max_agg and max_agg[0]:
                    max_dt = datetime.combine(max_agg[0], time.min)
                    if max_dt < cutoff:
                        cutoff = max_dt
        except Exception as e:
            log.warning("cleanup_old check daily_attendance: %s", e)
        cur.execute("DELETE FROM attendance WHERE event_time < %s", (cutoff,))
        log.info("Eski yozuvlar o'chirildi: %s", cur.rowcount)


def update_daily_realtime(cur, emp_no, d: date):
    """Event sodir bo'lishi bilan daily_attendance jadvalini real-vaqtda darhol yangilash"""
    try:
        cur.execute("""
            WITH stats AS (
                SELECT 
                    MIN(event_time) FILTER (WHERE event_type = 'checkIn') as f_in,
                    MAX(event_time) FILTER (WHERE event_type = 'checkOut') as l_out,
                    COUNT(*) as ev_cnt,
                    COALESCE((SELECT late_minutes FROM attendance 
                              WHERE emp_no = %s AND event_time::date = %s AND event_type = 'checkIn' 
                              ORDER BY event_time ASC LIMIT 1), 0) as l_min
                FROM attendance
                WHERE emp_no = %s AND event_time::date = %s
            )
            INSERT INTO daily_attendance (emp_no, day, first_in, last_out, late_minutes, worked_minutes, events_count, status)
            SELECT 
                %s, %s, f_in, l_out, l_min,
                CASE WHEN f_in IS NOT NULL AND l_out IS NOT NULL AND l_out > f_in 
                     THEN ROUND(EXTRACT(EPOCH FROM (l_out - f_in)) / 60)
                     ELSE NULL END,
                ev_cnt,
                CASE WHEN f_in IS NOT NULL 
                     THEN (CASE WHEN l_min > 0 THEN 'late' ELSE 'present' END)
                     ELSE 'absent' END
            FROM stats
            ON CONFLICT (emp_no, day) DO UPDATE SET
                first_in = EXCLUDED.first_in,
                last_out = EXCLUDED.last_out,
                late_minutes = EXCLUDED.late_minutes,
                worked_minutes = EXCLUDED.worked_minutes,
                events_count = EXCLUDED.events_count,
                status = CASE 
                    WHEN daily_attendance.status IN ('leave', 'holiday', 'weekend') AND EXCLUDED.first_in IS NULL 
                    THEN daily_attendance.status 
                    ELSE EXCLUDED.status 
                END;
        """, (emp_no, d, emp_no, d, emp_no, d))
    except Exception as e:
        log.warning("Realtime daily update ogohlantirish: %s", e)


def save_event(emp_no, name, etype, t: datetime, late, serial_no=None):
    with db() as c, c.cursor() as cur:
        cur.execute("""INSERT INTO attendance (emp_no, employee_name, event_type, event_time, late_minutes, serial_no)
                       VALUES (%s,%s,%s,%s,%s,%s)""", (emp_no, name, etype, t, late, serial_no))
        update_daily_realtime(cur, emp_no, t.date())


_ci = {"day": None, "ids": set()}
_ci_lock = threading.Lock()


def has_checkin_today(emp_no, d: date) -> bool:
    """Kunda bitta marta DB dan o'qiladi, qolgan tekshiruvlar xotiradan."""
    with _ci_lock:
        if _ci["day"] != d:
            s0 = datetime.combine(d, time.min)
            with db() as c, c.cursor() as cur:
                cur.execute("""SELECT DISTINCT emp_no FROM attendance WHERE event_type='checkIn'
                               AND event_time >= %s AND event_time < %s""", (s0, s0 + timedelta(days=1)))
                _ci["ids"] = {r[0] for r in cur.fetchall()}
            _ci["day"] = d
        return emp_no in _ci["ids"]


def mark_checkin(emp_no, d: date):
    with _ci_lock:
        if _ci["day"] == d:
            _ci["ids"].add(emp_no)


def calc_late(t: datetime, start: time) -> int:
    base = datetime.combine(t.date(), start)
    if t <= base + timedelta(minutes=GRACE_MINUTES):
        return 0
    return max(1, int((t - base).total_seconds() // 60))


# ===================== TELEGRAMGA RASM YUBORISH (navbat) =====================
SEND_Q = queue.Queue(maxsize=50)
UNKNOWN_SEEN = {}  # ro'yxatda yo'q, lekin qurilmadan kelgan ID lar (xotirada)
_last_event = {}   # (emp_no, type) -> datetime  (takroriy signalni filtrlash)


def tg_targets():
    with LOCK:
        groups = list(CHATS)
    targets = list(dict.fromkeys(TG_CHAT_IDS + [str(g) for g in groups]))
    if not targets:   # chat ham, guruh ham yo'q bo'lsa — adminlarning shaxsiy chatiga
        with LOCK:
            targets = [str(a) for a in (SUPER_ADMINS | ADMINS)]
        if not targets:
            log.warning("Xabar yuboradigan chat yo'q: TELEGRAM_CHAT_ID yoki ADMIN_IDS ni to'ldiring")
    return targets


def tg_worker():
    """Bitta fon oqimi navbatdan oladi: server ustiga thread'lar to'planmaydi.
    Rasm faqat bir marta yuklanadi, qolgan chatlarga file_id orqali yuboriladi."""
    session = requests.Session()
    url = f"https://api.telegram.org/bot{TG_TOKEN}/sendPhoto"
    while True:
        caption, img = SEND_Q.get()
        try:
            file_id = None
            for chat in tg_targets():
                try:
                    data = {"chat_id": chat, "caption": caption, "parse_mode": "HTML"}
                    if file_id:
                        data["photo"] = file_id
                        r = session.post(url, data=data, timeout=15)
                    else:
                        r = session.post(url, data=data, files={"photo": ("face.jpg", img)}, timeout=25)
                    j = r.json()
                    if j.get("ok"):
                        log.debug("  -> Telegramga yuborildi: %s", chat)
                        if not file_id:
                            file_id = j["result"]["photo"][-1]["file_id"]
                    else:
                        log.warning("Telegram xato (%s): %s", chat, j.get("description"))
                except Exception as e:
                    log.warning("Yuborishda xato (%s): %s", chat, e)
        finally:
            SEND_Q.task_done()


def enqueue_photo(caption, img):
    try:
        SEND_Q.put_nowait((caption, img))
    except queue.Full:
        log.warning("Navbat to'lgan, xabar tashlab yuborildi")


# ===================== HIKVISION WEBHOOK =====================
def extract_image(inner, data):
    if request.files:
        for key in request.files:
            f = request.files[key]
            if f.filename:
                return f.read()
    b64 = inner.get("facePicBinaryData") or data.get("facePicBinaryData")
    if b64:
        try:
            return base64.b64decode(b64)
        except Exception:
            pass
    pic_url = inner.get("pictureURL") or data.get("pictureURL")
    if pic_url:
        if pic_url.startswith("/"):
            pic_url = f"http://{HIK_IP}{pic_url}"
        try:
            r = requests.get(pic_url, auth=HTTPDigestAuth(HIK_USER, HIK_PASS), timeout=5)
            if r.status_code == 200:
                return r.content
        except Exception:
            pass
    return None


def _trunc(o, n=120):
    if isinstance(o, dict):
        return {k: _trunc(v, n) for k, v in o.items()}
    if isinstance(o, list):
        return [_trunc(v, n) for v in o[:10]]
    if isinstance(o, str) and len(o) > n:
        return o[:n] + f"...({len(o)} belgi)"
    return o


@app.route("/", methods=["POST"])
def hikvision_listener():
    event_data = request.form.get("AccessControllerEvent")
    if not event_data:
        body = request.get_data(cache=True, as_text=False)[:600]
        log.debug("DIQQAT: AccessControllerEvent maydoni yo'q. content_type=%s form_keys=%s files=%s body=%r",
                 request.content_type, list(request.form.keys()), list(request.files.keys()), body)
        return "OK", 200
    # Eshik ochildi/yopildi kabi eventlarni JSON ochmasdan tashlaymiz
    if "attendanceStatus" not in event_data and '"label"' not in event_data:
        return "OK", 200
    try:
        data = json.loads(event_data)
        if data.get("eventType") == "heartBeat":
            return "OK", 200
        inner = data.get("AccessControllerEvent", {}) or {}
        if os.getenv("DEBUG_EVENTS", "0") == "1":
            log.info("RAW: %s", json.dumps(_trunc(data), ensure_ascii=False))

        emp_no = str(inner.get("employeeNoString") or "").strip()
        dev_name = inner.get("name")
        log.debug("EVENT: id=%r name=%r status=%r label=%r time=%r files=%s",
                 emp_no, dev_name, inner.get("attendanceStatus"), inner.get("label"),
                 data.get("dateTime"), list(request.files.keys()))
        if not emp_no and not dev_name:
            log.debug("  -> tashlandi: xodim tanilmagan (ID va ism yo'q)")
            return "OK", 200
        emp_no = emp_no or dev_name

        # Vaqt (server vaqt mintaqasidan qat'i nazar Toshkent vaqtiga keltiriladi)
        try:
            t = datetime.fromisoformat(data.get("dateTime", ""))
            if t.tzinfo:
                t = t.astimezone(TZ).replace(tzinfo=None)
        except Exception:
            t = now_local()

        attendance = inner.get("attendanceStatus")
        label = inner.get("label")
        if attendance == "checkIn" or label == "Keldim":
            etype = "checkIn"
        elif attendance == "checkOut" or label == "Ketdim":
            etype = "checkOut"
        else:
            log.debug("  -> tashlandi: Keldim/Ketdim emas (status/label mos kelmadi)")
            return "OK", 200

        # Faqat admin ro'yxatga qo'shgan, faol xodimlar
        with LOCK:
            info = EMP.get(emp_no)
        if not info or not info["active"]:
            log.debug("  -> tashlandi: ro'yxatda yo'q yoki nofaol (ID %s)", emp_no)
            if not info:
                UNKNOWN_SEEN[emp_no] = dev_name or ""
                if len(UNKNOWN_SEEN) > 50:
                    UNKNOWN_SEEN.pop(next(iter(UNKNOWN_SEEN)))
            return "OK", 200
        name = info["name"]

        # Rasm bo'lmasa hech narsa qilinmaydi (avvalgi mantiq saqlandi)
        image = extract_image(inner, data)
        if not image:
            log.debug("  -> tashlandi: rasm topilmadi")
            return "OK", 200

        # Takroriy signal (30 soniya ichida)
        key = (emp_no, etype)
        prev = _last_event.get(key)
        if prev and abs((t - prev).total_seconds()) < 30:
            log.debug("  -> tashlandi: 30 soniya ichida takror")
            return "OK", 200
        if len(_last_event) > 500:
            _last_event.clear()
        _last_event[key] = t

        late = 0
        if etype == "checkIn":
            if has_checkin_today(emp_no, t.date()):
                status = "Qayta o'tish 🔁"
            else:
                late = calc_late(t, info["start"])
                status = "Kech keldi ⚠️" if late > 0 else "Vaqtida keldi ✅"
        else:
            status = "Ketdi 🚶‍♂️"

        serial_no = inner.get("serialNo")
        try:
            serial_no = int(serial_no) if serial_no is not None else None
        except Exception:
            serial_no = None

        try:
            save_event(emp_no, name, etype, t, late, serial_no=serial_no)
        except Exception as e:
            log.error("Bazaga yozishda xato: %s", e)

        caption = (f"🚪 {status}\n👤 Ism-familiya: {esc(name)}\n"
                   f"📅 Sana: {t:%d.%m.%Y}\n🕐 Vaqt: {t:%H:%M:%S}")
        if etype == "checkIn":
            caption += f"\n⏰ Ish boshlanishi: {info['start']:%H:%M}"
        if late > 0:
            caption += f"\n⚠️ Kechikish: {fmt_min(late)}"
        if etype == "checkIn":
            mark_checkin(emp_no, t.date())
        log.info("✔ %s | %s | %s | kechikish=%s", name, etype, f"{t:%H:%M:%S}", late)
        enqueue_photo(caption, image)
    except Exception as e:
        log.exception("Webhook xatosi: %s", e)
    return "OK", 200


# ===================== HISOBOTLAR =====================
def today_text() -> str:
    d = now_local().date()
    s = datetime.combine(d, time.min)
    with db() as c, c.cursor() as cur:
        cur.execute("""
            SELECT emp_no, MAX(employee_name),
                   MIN(event_time) FILTER (WHERE event_type='checkIn'),
                   MAX(event_time) FILTER (WHERE event_type='checkOut'),
                   COALESCE(MAX(late_minutes),0)
            FROM attendance WHERE event_time >= %s AND event_time < %s
            GROUP BY emp_no ORDER BY 3 NULLS LAST""", (s, s + timedelta(days=1)))
        rows = cur.fetchall()
    msg = f"📅 <b>Bugungi hisobot ({d:%d.%m.%Y})</b>\n\n"
    came = set()
    if not rows:
        msg += "Bugun hali hech kim o'tmadi.\n"
    for emp_no, name, cin, cout, late in rows:
        came.add(emp_no)
        line = f"👤 {esc(name)}: "
        if cin:
            line += f"{'⚠️' if late > 0 else '✅'} {cin:%H:%M}"
            if late > 0:
                line += f" (+{fmt_min(late)})"
        if cout:
            line += f"  🚶 {cout:%H:%M}"
        msg += line + "\n"
    with LOCK:
        absent = [v["name"] for k, v in EMP.items() if v["active"] and k not in came]
    if absent and rows:
        msg += "\n❌ <b>Hali kelmaganlar:</b> " + ", ".join(esc(n) for n in sorted(absent))
    return msg


def weekly_text() -> str:
    end = now_local()
    start = datetime.combine((end - timedelta(days=6)).date(), time.min)
    with db() as c, c.cursor() as cur:
        cur.execute("""
            SELECT MAX(employee_name),
                   COUNT(DISTINCT event_time::date),
                   COUNT(DISTINCT event_time::date) FILTER (WHERE late_minutes > 0),
                   COALESCE(SUM(late_minutes),0)
            FROM attendance WHERE event_type='checkIn' AND event_time >= %s
            GROUP BY emp_no ORDER BY 4 DESC""", (start,))
        rows = cur.fetchall()
    msg = "📊 <b>Haftalik hisobot (oxirgi 7 kun)</b>\n\n"
    if not rows:
        return msg + "Ma'lumot yo'q."
    for name, days, late_days, late_sum in rows:
        msg += f"👤 {esc(name)}: {days} kun keldi, {late_days} kun kech qoldi, jami {fmt_min(late_sum)}\n"
    return msg


DAILY_SQL = """
    SELECT emp_no, MAX(employee_name), event_time::date,
           MIN(event_time) FILTER (WHERE event_type='checkIn'),
           MAX(event_time) FILTER (WHERE event_type='checkOut'),
           COALESCE(MAX(late_minutes),0)
    FROM attendance WHERE event_time >= %s AND event_time < %s {flt}
    GROUP BY emp_no, event_time::date ORDER BY emp_no, 3"""


def worked(cin, cout) -> str:
    if cin and cout and cout > cin:
        m = int((cout - cin).total_seconds() // 60)
        return f"{m // 60}:{m % 60:02d}"
    return ""


def style_sheet(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(w, 10), 40)
    ws.freeze_panes = "A2"


def workdays(start: date, end: date):
    today = now_local().date()
    d = start
    while d <= min(end, today):
        if d.weekday() not in WEEKEND_DAYS:
            yield d
        d += timedelta(days=1)


def build_excel(start: date, end: date, emp_no=None):
    s = datetime.combine(start, time.min)
    e = datetime.combine(end, time.min) + timedelta(days=1)
    with db() as c, c.cursor() as cur:
        if emp_no:
            cur.execute(DAILY_SQL.format(flt="AND emp_no = %s"), (s, e, emp_no))
        else:
            cur.execute(DAILY_SQL.format(flt=""), (s, e))
        rows = cur.fetchall()
        events = []
        if emp_no:
            cur.execute("""SELECT event_time, event_type, late_minutes FROM attendance
                           WHERE emp_no=%s AND event_time >= %s AND event_time < %s
                           ORDER BY event_time""", (emp_no, s, e))
            events = cur.fetchall()

    wb = Workbook()
    ws = wb.active
    period = f"{start:%d.%m.%Y}-{end:%d.%m.%Y}"

    if emp_no:
        with LOCK:
            info = EMP.get(emp_no, {"name": emp_no, "start": DEFAULT_START})
        by_day = {r[2]: r for r in rows}
        ws.title = "Hisobot"
        ws.append(["Sana", "Hafta kuni", "Keldi", "Ketdi", "Ish vaqti", "Kechikish (daq)", "Holat"])
        came = late_days = late_sum = absent = 0
        today = now_local().date()
        d = start
        while d <= end:
            r = by_day.get(d)
            wd = WEEKDAYS[d.weekday()]
            if r and r[3]:
                came += 1
                late = r[5]
                if late > 0:
                    late_days += 1
                    late_sum += late
                ws.append([f"{d:%d.%m.%Y}", wd, f"{r[3]:%H:%M}", f"{r[4]:%H:%M}" if r[4] else "",
                           worked(r[3], r[4]), late, "Kech keldi" if late > 0 else "Vaqtida"])
            elif d.weekday() in WEEKEND_DAYS:
                ws.append([f"{d:%d.%m.%Y}", wd, "", "", "", "", "Dam olish"])
            elif d <= today:
                absent += 1
                ws.append([f"{d:%d.%m.%Y}", wd, "", "", "", "", "Kelmadi"])
            d += timedelta(days=1)
        ws.append([])
        ws.append(["JAMI", "", f"Keldi: {came} kun", f"Kelmadi: {absent} kun", "",
                   late_sum, f"Kechikkan: {late_days} kun"])
        ws.append([f"Xodim: {info['name']} (ID {emp_no}), ish boshlanishi: {info['start']:%H:%M}"])
        ws2 = wb.create_sheet("Barcha o'tishlar")
        ws2.append(["Vaqt", "Turi", "Kechikish (daq)"])
        names = {"checkIn": "Keldi", "checkOut": "Ketdi"}
        for t, typ, late in events:
            ws2.append([f"{t:%d.%m.%Y %H:%M:%S}", names.get(typ, typ), late or 0])
        style_sheet(ws2)
        style_sheet(ws)
        fname = f"hisobot_{re.sub(r'[^A-Za-z0-9_-]', '_', info['name'])}_{period}.xlsx"
    else:
        ws.title = "Umumiy"
        ws.append(["Xodim", "Ish boshlanishi", "Kelgan kunlar", "Kelmagan kunlar",
                   "Kechikkan kunlar", "Jami kechikish (daq)", "O'rtacha kechikish (daq)"])
        det = wb.create_sheet("Kunlik")
        det.append(["Xodim", "Sana", "Hafta kuni", "Keldi", "Ketdi", "Ish vaqti", "Kechikish (daq)"])
        stats = {}
        for no, name, d, cin, cout, late in rows:
            st = stats.setdefault(no, {"name": name, "came": 0, "late_days": 0, "late_sum": 0})
            if cin:
                st["came"] += 1
                if late > 0:
                    st["late_days"] += 1
                    st["late_sum"] += late
            det.append([name, f"{d:%d.%m.%Y}", WEEKDAYS[d.weekday()],
                        f"{cin:%H:%M}" if cin else "", f"{cout:%H:%M}" if cout else "",
                        worked(cin, cout), late])
        with LOCK:
            for no, v in EMP.items():
                if v["active"] and no not in stats:
                    stats[no] = {"name": v["name"], "came": 0, "late_days": 0, "late_sum": 0}
            starts = {no: v["start"] for no, v in EMP.items()}
        wd_total = sum(1 for _ in workdays(start, end))
        for no, st in sorted(stats.items(), key=lambda x: x[1]["name"]):
            avg = round(st["late_sum"] / st["late_days"]) if st["late_days"] else 0
            ws.append([st["name"], f"{starts.get(no, DEFAULT_START):%H:%M}", st["came"],
                       max(wd_total - st["came"], 0), st["late_days"], st["late_sum"], avg])
        style_sheet(ws)
        style_sheet(det)
        fname = f"hisobot_barcha_{period}.xlsx"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, fname


# ===================== TELEGRAM BOT =====================
B_TODAY, B_WEEK, B_EXCEL = "📅 Bugungi hisobot", "📊 Haftalik hisobot", "📥 Excel hisobot"
B_EMPS, B_ADD_EMP, B_ADMIN = "👥 Xodimlar", "➕ Xodim qo'shish", "⚙️ Admin panel"
MENU_BUTTONS = {B_TODAY, B_WEEK, B_EXCEL, B_EMPS, B_ADD_EMP, B_ADMIN}
MAIN_KB = ReplyKeyboardMarkup(
    [[B_TODAY, B_WEEK], [B_EXCEL, B_EMPS], [B_ADD_EMP, B_ADMIN]], resize_keyboard=True)
XL_SEM = asyncio.Semaphore(1)   # bir vaqtda faqat bitta Excel tayyorlanadi


def ikb(rows):
    return InlineKeyboardMarkup([[InlineKeyboardButton(t, callback_data=d) for t, d in r] for r in rows])


async def reply_long(msg, text, **kw):
    chunk = ""
    for line in text.split("\n"):
        if len(chunk) + len(line) + 1 > 3800:
            await msg.reply_text(chunk, parse_mode="HTML", **kw)
            chunk = ""
        chunk += line + "\n"
    if chunk.strip():
        await msg.reply_text(chunk, parse_mode="HTML", **kw)


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    if not is_admin(uid):
        await update.message.reply_text(
            f"⛔ Sizda ruxsat yo'q.\nSizning Telegram ID: {uid}\nShu ID ni admin'ga yuboring.")
        return
    context.user_data.clear()
    await update.message.reply_text("Xush kelibsiz! Kerakli bo'limni tanlang:", reply_markup=MAIN_KB)


def period_range(p: str):
    today = now_local().date()
    if p == "1":
        return today, today
    if p == "7":
        return today - timedelta(days=6), today
    if p == "30":
        return today - timedelta(days=29), today
    if p == "m0":
        return today.replace(day=1), today
    if p == "m1":
        last = today.replace(day=1) - timedelta(days=1)
        return last.replace(day=1), last
    return None


def period_kb(scope):
    return ikb([
        [("Bugun", f"per:{scope}:1"), ("7 kun", f"per:{scope}:7")],
        [("30 kun", f"per:{scope}:30"), ("Shu oy", f"per:{scope}:m0")],
        [("O'tgan oy", f"per:{scope}:m1"), ("📝 Sana kiritish", f"per:{scope}:c")],
    ])


def parse_dates(text):
    found = re.findall(r"\d{1,4}[./-]\d{1,2}[./-]\d{1,4}", text)
    out = []
    for f in found[:2]:
        for fmt in ("%d.%m.%Y", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y.%m.%d"):
            try:
                out.append(datetime.strptime(f, fmt).date())
                break
            except ValueError:
                continue
    if len(out) == 1:
        out.append(out[0])
    if len(out) != 2:
        return None
    return (out[0], out[1]) if out[0] <= out[1] else (out[1], out[0])


async def send_excel(context, chat_id, scope, start_d, end_d):
    if (end_d - start_d).days > 366:
        await context.bot.send_message(chat_id, "❌ Muddat 1 yildan oshmasligi kerak.")
        return
    async with XL_SEM:
        emp_no = None if scope == "all" else scope
        buf, fname = await asyncio.to_thread(build_excel, start_d, end_d, emp_no)
    await context.bot.send_document(chat_id, document=buf, filename=fname,
                                    caption=f"📥 {start_d:%d.%m.%Y} — {end_d:%d.%m.%Y}")


async def show_employees(msg):
    with LOCK:
        items = sorted(((k, v) for k, v in EMP.items() if v["active"]), key=lambda x: x[1]["name"])
    if not items:
        await msg.reply_text("Hozircha xodimlar yo'q. «➕ Xodim qo'shish» ni bosing.")
        return
    rows, row = [], []
    for k, v in items[:60]:
        row.append((f"{v['name']} • {v['start']:%H:%M}", f"emp:{k}"))
        if len(row) == 2:
            rows.append(row); row = []
    if row:
        rows.append(row)
    extra = f"\n(Dastlabki 60 tasi ko'rsatildi, jami {len(items)})" if len(items) > 60 else ""
    await msg.reply_text("👥 Xodimni tanlang:" + extra, reply_markup=ikb(rows))


async def on_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    uid = update.effective_user.id
    if not is_admin(uid):
        await update.message.reply_text(f"⛔ Ruxsat yo'q. Sizning ID: {uid}")
        return
    text = (update.message.text or "").strip()
    ud = context.user_data
    msg = update.message

    if text in MENU_BUTTONS:
        ud.pop("state", None)
        if text == B_TODAY:
            await reply_long(msg, await asyncio.to_thread(today_text))
        elif text == B_WEEK:
            t = await asyncio.to_thread(weekly_text)
            await msg.reply_text(t, parse_mode="HTML", reply_markup=ikb([[("🔄 Yangilash", "refresh_weekly")]]))
        elif text == B_EXCEL:
            await msg.reply_text("📥 Qaysi hisobot kerak?", reply_markup=ikb([
                [("👥 Barcha xodimlar", "xl_all"), ("👤 Bitta xodim", "xl_one")]]))
        elif text == B_EMPS:
            await show_employees(msg)
        elif text == B_ADD_EMP:
            ud["state"] = "add_emp"
            await msg.reply_text(
                "Xodim ma'lumotini shu formatda yuboring:\n\n"
                "<code>ID | Ism Familiya | Ish boshlanishi</code>\n"
                "Masalan: <code>15 | Aliyev Vali | 8:30</code>\n\n"
                "ID — Hikvision qurilmasidagi xodim raqami. Vaqt yozilmasa "
                f"{DEFAULT_START:%H:%M} qo'yiladi.", parse_mode="HTML")
        elif text == B_ADMIN:
            await msg.reply_text("⚙️ Admin panel", reply_markup=ikb([
                [("➕ Admin qo'shish", "adm_add"), ("➖ Admin o'chirish", "adm_del")],
                [("📋 Adminlar ro'yxati", "adm_list")]]))
        return

    state = ud.get("state")
    if state == "add_emp":
        parts = [p.strip() for p in text.split("|")]
        if len(parts) < 2 or not parts[0] or not parts[1]:
            await msg.reply_text("❌ Format noto'g'ri. Masalan: 15 | Aliyev Vali | 8:30")
            return
        start_t = parse_hhmm(parts[2]) if len(parts) > 2 else DEFAULT_START
        if not start_t:
            await msg.reply_text("❌ Vaqt noto'g'ri. Masalan: 8:30")
            return
        await asyncio.to_thread(upsert_employee, parts[0], parts[1], start_t)
        ud.pop("state", None)
        await msg.reply_text(f"✅ Qo'shildi: {parts[1]} (ID {parts[0]}), ish boshlanishi {start_t:%H:%M}")
    elif state == "set_time":
        t = parse_hhmm(text)
        if not t:
            await msg.reply_text("❌ Vaqt noto'g'ri. Masalan: 8 yoki 8:30")
            return
        emp_no = ud.pop("emp_no", None)
        ud.pop("state", None)
        await asyncio.to_thread(set_work_start, emp_no, t)
        await msg.reply_text(f"✅ Ish boshlanish vaqti {t:%H:%M} qilib belgilandi.")
    elif state == "add_admin":
        parts = [p.strip() for p in text.split("|")]
        if not parts[0].lstrip("-").isdigit():
            await msg.reply_text("❌ Faqat raqamli Telegram ID yuboring. Masalan: 123456789 | Ism")
            return
        await asyncio.to_thread(add_admin, int(parts[0]), parts[1] if len(parts) > 1 else "")
        ud.pop("state", None)
        await msg.reply_text("✅ Admin qo'shildi.")
    elif state == "xl_range":
        rng = parse_dates(text)
        if not rng:
            await msg.reply_text("❌ Sana noto'g'ri. Masalan: 01.10.2026 - 22.10.2026")
            return
        scope = ud.pop("xl_scope", "all")
        ud.pop("state", None)
        await msg.reply_text("⏳ Hisobot tayyorlanmoqda...")
        await send_excel(context, msg.chat_id, scope, *rng)
    else:
        await msg.reply_text("Menyudan tanlang 👇", reply_markup=MAIN_KB)


async def on_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    if not is_admin(q.from_user.id):
        await q.answer("Ruxsat yo'q", show_alert=True)
        return
    await q.answer()
    data, ud = q.data, context.user_data
    chat_id = q.message.chat_id

    if data == "refresh_weekly":
        t = await asyncio.to_thread(weekly_text)
        try:
            await q.edit_message_text(t + f"\n🕐 {now_local():%H:%M:%S}", parse_mode="HTML",
                                      reply_markup=ikb([[("🔄 Yangilash", "refresh_weekly")]]))
        except BadRequest:
            pass

    # ---- Excel ----
    elif data == "xl_all":
        await q.edit_message_text("📅 Muddatni tanlang (barcha xodimlar):", reply_markup=period_kb("all"))
    elif data == "xl_one":
        with LOCK:
            items = sorted(((k, v) for k, v in EMP.items() if v["active"]), key=lambda x: x[1]["name"])
        if not items:
            await q.edit_message_text("Xodimlar yo'q.")
            return
        rows, row = [], []
        for k, v in items[:60]:
            row.append((v["name"], f"xlp:{k}"))
            if len(row) == 2:
                rows.append(row); row = []
        if row:
            rows.append(row)
        await q.edit_message_text("👤 Xodimni tanlang:", reply_markup=ikb(rows))
    elif data.startswith("xlp:"):
        await q.edit_message_text("📅 Muddatni tanlang:", reply_markup=period_kb(data[4:]))
    elif data.startswith("per:"):
        _, scope, p = data.split(":", 2)
        if p == "c":
            ud["state"], ud["xl_scope"] = "xl_range", scope
            await q.edit_message_text("📝 Sanalarni yuboring. Masalan:\n01.10.2026 - 22.10.2026")
            return
        rng = period_range(p)
        await q.edit_message_text("⏳ Hisobot tayyorlanmoqda...")
        await send_excel(context, chat_id, scope, *rng)

    # ---- Xodim kartasi ----
    elif data.startswith("emp:"):
        no = data[4:]
        with LOCK:
            v = EMP.get(no)
        if not v:
            await q.edit_message_text("Topilmadi.")
            return
        await q.edit_message_text(
            f"👤 {v['name']}\n🆔 ID: {no}\n⏰ Ish boshlanishi: {v['start']:%H:%M}",
            reply_markup=ikb([[("⏰ Vaqtni o'zgartirish", f"wt:{no}")],
                              [("📥 Excel", f"xlp:{no}"), ("🗑 O'chirish", f"del:{no}")]]))
    elif data.startswith("wt:"):
        ud["state"], ud["emp_no"] = "set_time", data[3:]
        await q.edit_message_text("⏰ Yangi ish boshlanish vaqtini yuboring (masalan: 8 yoki 8:30):")
    elif data.startswith("del:"):
        await asyncio.to_thread(set_active, data[4:], False)
        await q.edit_message_text("🗑 Xodim o'chirildi (hisobotlarda va xabarlarda ko'rinmaydi).")

    # ---- Adminlar ----
    elif data == "adm_add":
        ud["state"] = "add_admin"
        await q.edit_message_text(
            "Yangi adminning Telegram ID sini yuboring:\n<code>123456789 | Ism</code>\n\n"
            "ID ni bilish uchun u botga /start yozsa, bot ID sini ko'rsatadi.", parse_mode="HTML")
    elif data == "adm_list":
        rows = await asyncio.to_thread(list_admins)
        t = "👮 <b>Adminlar:</b>\n" + "\n".join(f"• {i} {esc(n)}" for i, n in rows)
        if SUPER_ADMINS:
            t += "\n\n🔒 Asosiy: " + ", ".join(map(str, SUPER_ADMINS))
        await q.edit_message_text(t, parse_mode="HTML")
    elif data == "adm_del":
        rows = await asyncio.to_thread(list_admins)
        if not rows:
            await q.edit_message_text("O'chiriladigan admin yo'q.")
            return
        await q.edit_message_text("Qaysi adminni o'chirasiz?", reply_markup=ikb(
            [[(f"{n or i} ({i})", f"rma:{i}")] for i, n in rows]))
    elif data.startswith("rma:"):
        await asyncio.to_thread(remove_admin, int(data[4:]))
        await q.edit_message_text("✅ Admin o'chirildi.")


async def cmd_new(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update.effective_user.id):
        return
    if not UNKNOWN_SEEN:
        await update.message.reply_text("Ro'yxatda yo'q yangi ID kelmagan.")
        return
    t = "🆕 Qurilmadan kelgan, lekin ro'yxatda yo'q ID lar:\n\n"
    t += "\n".join(f"<code>{esc(k)} | {esc(v or 'Ism')}</code>" for k, v in UNKNOWN_SEEN.items())
    t += "\n\nQo'shish uchun «➕ Xodim qo'shish» ni bosib, shu qatorni (kerak bo'lsa vaqt bilan) yuboring."
    await reply_long(update.message, t)


async def on_my_chat_member(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Bot guruhga qo'shilsa — kelganlar rasmi shu guruhga ham boradi.
    Faqat admin qo'shgan guruhlar qabul qilinadi."""
    cm = update.my_chat_member
    chat = cm.chat
    if chat.type not in ("group", "supergroup"):
        return
    status = cm.new_chat_member.status
    if status in ("member", "administrator"):
        if not is_admin(cm.from_user.id):
            await context.bot.leave_chat(chat.id)
            return
        await asyncio.to_thread(add_chat, chat.id, chat.title)
        await context.bot.send_message(chat.id, "✅ Bot ulandi. Xodimlar davomati shu guruhga yuboriladi.")
    elif status in ("left", "kicked"):
        await asyncio.to_thread(remove_chat, chat.id)


async def check_config_job(context: ContextTypes.DEFAULT_TYPE):
    await asyncio.to_thread(check_config_version)


async def cleanup_job(context: ContextTypes.DEFAULT_TYPE):
    await asyncio.to_thread(cleanup_old)


async def on_error(update, context: ContextTypes.DEFAULT_TYPE):
    log.error("Bot xatosi: %s", context.error)


# ===================== ISHGA TUSHIRISH =====================
def run_flask():
    try:
        from waitress import serve
        serve(app, host="0.0.0.0", port=PORT, threads=2)
    except ImportError:
        app.run(host="0.0.0.0", port=PORT, debug=False, use_reloader=False, threaded=True)


if __name__ == "__main__":
    init_db()
    load_caches()
    threading.Thread(target=tg_worker, daemon=True).start()
    threading.Thread(target=run_flask, daemon=True).start()

    application = Application.builder().token(TG_TOKEN).build()
    private = filters.ChatType.PRIVATE
    application.add_handler(CommandHandler("start", start, filters=private))
    application.add_handler(CommandHandler("yangi", cmd_new, filters=private))
    application.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND & private, on_text))
    application.add_handler(CallbackQueryHandler(on_callback))
    application.add_handler(ChatMemberHandler(on_my_chat_member, ChatMemberHandler.MY_CHAT_MEMBER))
    application.add_error_handler(on_error)
    if application.job_queue:
        application.job_queue.run_daily(cleanup_job, time=time(3, 0, tzinfo=TZ))
        application.job_queue.run_repeating(check_config_job, interval=30, first=15)

    log.info("🤖 Bot va server ishga tushdi...")
    application.run_polling(allowed_updates=Update.ALL_TYPES, drop_pending_updates=True)
