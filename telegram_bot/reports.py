import io
import re
from datetime import datetime, date, time, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from database import db
from config import now_local, esc, fmt_min, WEEKDAYS, WEEKEND_DAYS, DEFAULT_START
from cache import EMP, LOCK

DAILY_SQL = """
    SELECT emp_no, MAX(employee_name), event_time::date,
           MIN(event_time) FILTER (WHERE event_type='checkIn'),
           MAX(event_time) FILTER (WHERE event_type='checkOut'),
           COALESCE(MAX(late_minutes),0)
    FROM attendance WHERE event_time >= %s AND event_time < %s {flt}
    GROUP BY emp_no, event_time::date ORDER BY emp_no, 3"""


def today_text() -> str:
    """Bugungi kun bo'yicha qisqacha matnli hisobot"""
    d = now_local().date()
    s = datetime.combine(d, time.min)
    with db() as c, c.cursor() as cur:
        cur.execute("""
            SELECT emp_no, MAX(employee_name),
                   MIN(event_time) FILTER (WHERE event_type='checkIn'),
                   MAX(event_time) FILTER (WHERE event_type='checkOut'),
                   COALESCE(MAX(late_minutes),0)
            FROM attendance WHERE event_time >= %s AND event_time < %s
            GROUP BY emp_no ORDER BY 3 NULLS LAST
        """, (s, s + timedelta(days=1)))
        rows = cur.fetchall()

    msg = f"📅 <b>Bugungi hisobot ({d:%d.%m.%Y})</b>\n\n"
    came = set()
    if not rows:
        msg += "Bugun hali hech kim o'tmadi.\n"

    for emp_no, name, cin, cout, late in rows:
        came.add(emp_no)
        line = f"👤 {esc(name)}: "
        if cin:
            line += f"{'⚠️' if late > 0 else '✅'} {cin:%H:%M}"
            if late > 0:
                line += f" (+{fmt_min(late)})"
        if cout:
            line += f"  🚶 {cout:%H:%M}"
        msg += line + "\n"

    with LOCK:
        absent = [v["name"] for k, v in EMP.items() if v["active"] and k not in came]

    if absent and rows:
        msg += "\n❌ <b>Hali kelmaganlar:</b> " + ", ".join(esc(n) for n in sorted(absent))
    return msg


def weekly_text() -> str:
    """Oxirgi 7 kunlik kechikishlar va kelgan kunlar hisoboti"""
    end = now_local()
    start = datetime.combine((end - timedelta(days=6)).date(), time.min)
    with db() as c, c.cursor() as cur:
        cur.execute("""
            SELECT MAX(employee_name),
                   COUNT(DISTINCT event_time::date),
                   COUNT(DISTINCT event_time::date) FILTER (WHERE late_minutes > 0),
                   COALESCE(SUM(late_minutes),0)
            FROM attendance WHERE event_type='checkIn' AND event_time >= %s
            GROUP BY emp_no ORDER BY 4 DESC
        """, (start,))
        rows = cur.fetchall()

    msg = "📊 <b>Haftalik hisobot (oxirgi 7 kun)</b>\n\n"
    if not rows:
        return msg + "Ma'lumot yo'q."

    for name, days, late_days, late_sum in rows:
        msg += f"👤 {esc(name)}: {days} kun keldi, {late_days} kun kech qoldi, jami {fmt_min(late_sum)}\n"
    return msg


def worked(cin, cout) -> str:
    if cin and cout and cout > cin:
        m = int((cout - cin).total_seconds() // 60)
        return f"{m // 60}:{m % 60:02d}"
    return ""


def style_sheet(ws):
    """Excel varag'iga sarlavha va ustun kengliklarini berish"""
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(w, 10), 40)
    ws.freeze_panes = "A2"


def workdays(start: date, end: date):
    """Dam olish kunlaridan tashqari ish kunlari generatori"""
    today = now_local().date()
    d = start
    while d <= min(end, today):
        if d.weekday() not in WEEKEND_DAYS:
            yield d
        d += timedelta(days=1)


def build_excel(start: date, end: date, emp_no=None):
    """Excel faylini xotirada (BytesIO) yaratish"""
    s = datetime.combine(start, time.min)
    e = datetime.combine(end, time.min) + timedelta(days=1)
    with db() as c, c.cursor() as cur:
        if emp_no:
            cur.execute(DAILY_SQL.format(flt="AND emp_no = %s"), (s, e, emp_no))
        else:
            cur.execute(DAILY_SQL.format(flt=""), (s, e))
        rows = cur.fetchall()
        events = []
        if emp_no:
            cur.execute("""
                SELECT event_time, event_type, late_minutes FROM attendance
                WHERE emp_no=%s AND event_time >= %s AND event_time < %s
                ORDER BY event_time
            """, (emp_no, s, e))
            events = cur.fetchall()

    wb = Workbook()
    ws = wb.active
    period = f"{start:%d.%m.%Y}-{end:%d.%m.%Y}"

    if emp_no:
        with LOCK:
            info = EMP.get(emp_no, {"name": emp_no, "start": DEFAULT_START})
        by_day = {r[2]: r for r in rows}
        ws.title = "Hisobot"
        ws.append(["Sana", "Hafta kuni", "Keldi", "Ketdi", "Ish vaqti", "Kechikish (daq)", "Holat"])
        came = late_days = late_sum = absent = 0
        today = now_local().date()
        d = start
        while d <= end:
            r = by_day.get(d)
            wd = WEEKDAYS[d.weekday()]
            if r and r[3]:
                came += 1
                late = r[5]
                if late > 0:
                    late_days += 1
                    late_sum += late
                ws.append([
                    f"{d:%d.%m.%Y}", wd, f"{r[3]:%H:%M}", f"{r[4]:%H:%M}" if r[4] else "",
                    worked(r[3], r[4]), late, "Kech keldi" if late > 0 else "Vaqtida"
                ])
            elif d.weekday() in WEEKEND_DAYS:
                ws.append([f"{d:%d.%m.%Y}", wd, "", "", "", "", "Dam olish"])
            elif d <= today:
                absent += 1
                ws.append([f"{d:%d.%m.%Y}", wd, "", "", "", "", "Kelmadi"])
            d += timedelta(days=1)

        ws.append([])
        ws.append(["JAMI", "", f"Keldi: {came} kun", f"Kelmadi: {absent} kun", "",
                   late_sum, f"Kechikkan: {late_days} kun"])
        ws.append([f"Xodim: {info['name']} (ID {emp_no}), ish boshlanishi: {info['start']:%H:%M}"])

        ws2 = wb.create_sheet("Barcha o'tishlar")
        ws2.append(["Vaqt", "Turi", "Kechikish (daq)"])
        names = {"checkIn": "Keldi", "checkOut": "Ketdi"}
        for t, typ, late in events:
            ws2.append([f"{t:%d.%m.%Y %H:%M:%S}", names.get(typ, typ), late or 0])

        style_sheet(ws2)
        style_sheet(ws)
        safe_name = re.sub(r'[^A-Za-z0-9_-]', '_', info['name'])
        fname = f"hisobot_{safe_name}_{period}.xlsx"
    else:
        ws.title = "Umumiy"
        ws.append(["Xodim", "Ish boshlanishi", "Kelgan kunlar", "Kelmagan kunlar",
                   "Kechikkan kunlar", "Jami kechikish (daq)", "O'rtacha kechikish (daq)"])
        det = wb.create_sheet("Kunlik")
        det.append(["Xodim", "Sana", "Hafta kuni", "Keldi", "Ketdi", "Ish vaqti", "Kechikish (daq)"])
        stats = {}
        for no, name, d, cin, cout, late in rows:
            st = stats.setdefault(no, {"name": name, "came": 0, "late_days": 0, "late_sum": 0})
            if cin:
                st["came"] += 1
                if late > 0:
                    st["late_days"] += 1
                    st["late_sum"] += late
            det.append([
                name, f"{d:%d.%m.%Y}", WEEKDAYS[d.weekday()],
                f"{cin:%H:%M}" if cin else "", f"{cout:%H:%M}" if cout else "",
                worked(cin, cout), late
            ])

        with LOCK:
            for no, v in EMP.items():
                if v["active"] and no not in stats:
                    stats[no] = {"name": v["name"], "came": 0, "late_days": 0, "late_sum": 0}
            starts = {no: v["start"] for no, v in EMP.items()}

        wd_total = sum(1 for _ in workdays(start, end))
        for no, st in sorted(stats.items(), key=lambda x: x[1]["name"]):
            avg = round(st["late_sum"] / st["late_days"]) if st["late_days"] else 0
            ws.append([
                st["name"], f"{starts.get(no, DEFAULT_START):%H:%M}", st["came"],
                max(wd_total - st["came"], 0), st["late_days"], st["late_sum"], avg
            ])
        style_sheet(ws)
        style_sheet(det)
        fname = f"hisobot_barcha_{period}.xlsx"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, fname
