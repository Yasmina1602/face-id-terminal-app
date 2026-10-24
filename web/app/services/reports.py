import io
import re
import threading
from datetime import date, datetime, time, timedelta
from typing import Optional, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.models import Employee, DailyAttendance, AttendanceAdjustment, Attendance

WEEKDAYS = ["Dushanba", "Seshanba", "Chorshanba", "Payshanba", "Juma", "Shanba", "Yakshanba"]
XL_LOCK = threading.Lock()  # Bir vaqtda faqat bitta Excel eksport


def fmt_min(m: Optional[int]) -> str:
    m = int(m or 0)
    return f"{m} daq" if m < 60 else f"{m // 60} soat {m % 60} daq"


def worked_str(minutes: Optional[int]) -> str:
    if not minutes:
        return ""
    return f"{minutes // 60}:{minutes % 60:02d}"


def style_sheet(ws):
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    center = Alignment(horizontal="center", vertical="center")

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    ws.freeze_panes = "A2"


def export_employee_report(
    db: Session,
    emp_no: str,
    date_from: date,
    date_to: date
) -> Tuple[io.BytesIO, str]:
    with XL_LOCK:
        emp = db.query(Employee).filter(Employee.emp_no == emp_no).first()
        emp_name = emp.name if emp else emp_no
        work_start = emp.work_start if emp else time(9, 0)

        daily_records = (
            db.query(DailyAttendance)
            .filter(
                DailyAttendance.emp_no == emp_no,
                DailyAttendance.day >= date_from,
                DailyAttendance.day <= date_to
            )
            .all()
        )
        daily_map = {r.day: r for r in daily_records}

        adjustments = (
            db.query(AttendanceAdjustment)
            .filter(
                AttendanceAdjustment.emp_no == emp_no,
                AttendanceAdjustment.day >= date_from,
                AttendanceAdjustment.day <= date_to
            )
            .all()
        )
        adj_map = {a.day: a for a in adjustments}

        events = (
            db.query(Attendance)
            .filter(
                Attendance.emp_no == emp_no,
                Attendance.event_time >= datetime.combine(date_from, time.min),
                Attendance.event_time <= datetime.combine(date_to, time.max)
            )
            .order_by(Attendance.event_time.asc())
            .all()
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Davomat"

        # Headers
        ws.append([
            "Sana", "Hafta kuni", "Holat", "Kelgan", "Ketgan",
            "Ish vaqti", "Kechikish (daq)", "Sabab", "Jarima (so'm)"
        ])

        came_days = 0
        absent_days = 0
        late_days = 0
        total_late = 0
        total_fine = 0

        cur_day = date_from
        while cur_day <= date_to:
            wd_name = WEEKDAYS[cur_day.weekday()]
            rec = daily_map.get(cur_day)
            adj = adj_map.get(cur_day)

            first_in_str = f"{rec.first_in:%H:%M}" if (rec and rec.first_in) else ""
            last_out_str = f"{rec.last_out:%H:%M}" if (rec and rec.last_out) else ""
            worked_time = worked_str(rec.worked_minutes) if rec else ""
            late = rec.late_minutes if rec else 0
            status_text = rec.status if rec else "Noma'lum"

            reason = adj.reason if adj and adj.reason else ""
            fine = float(adj.fine_amount) if adj and adj.fine_amount else 0.0

            if rec and (rec.first_in or rec.status in ("present", "late")):
                came_days += 1
                if late > 0:
                    late_days += 1
                    total_late += late
                    status_text = "Kech keldi"
                else:
                    status_text = "Vaqtida keldi"
            elif rec and rec.status == "absent":
                absent_days += 1
                status_text = "Kelmadi"
            elif rec and rec.status == "weekend":
                status_text = "Dam olish kuni"
            elif rec and rec.status == "holiday":
                status_text = "Bayram"

            if adj and adj.status_override:
                status_text = f"{status_text} (Tuzatilgan: {adj.status_override})"

            total_fine += fine

            ws.append([
                f"{cur_day:%d.%m.%Y}",
                wd_name,
                status_text,
                first_in_str,
                last_out_str,
                worked_time,
                late if late > 0 else "",
                reason,
                fine if fine > 0 else ""
            ])
            cur_day += timedelta(days=1)

        ws.append([])
        # JAMI qatori
        ws.append([
            "JAMI", "",
            f"Keldi: {came_days} kun",
            f"Kelmadi: {absent_days} kun",
            f"Kechikkan: {late_days} kun",
            "",
            total_late,
            "",
            total_fine
        ])
        ws.append([f"Xodim: {emp_name} (ID: {emp_no}), Ish boshlanishi: {work_start:%H:%M}"])

        # Ikkinchi varaq: Barcha o'tishlar logi
        ws2 = wb.create_sheet("Barcha o'tishlar")
        ws2.append(["Vaqt", "Turi", "Kechikish (daq)", "Serial No"])
        for ev in events:
            ev_type_str = "Keldi" if ev.event_type == "checkIn" else "Ketdi"
            ws2.append([
                f"{ev.event_time:%d.%m.%Y %H:%M:%S}",
                ev_type_str,
                ev.late_minutes or 0,
                ev.serial_no or ""
            ])

        style_sheet(ws)
        style_sheet(ws2)

        safe_name = re.sub(r'[^A-Za-z0-9_-]', '_', emp_name)
        period = f"{date_from:%d.%m.%Y}_{date_to:%d.%m.%Y}"
        filename = f"davomat_{safe_name}_{period}.xlsx"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf, filename


def export_summary_report(
    db: Session,
    date_from: date,
    date_to: date
) -> Tuple[io.BytesIO, str]:
    with XL_LOCK:
        employees = db.query(Employee).order_by(Employee.name.asc()).all()
        daily_records = (
            db.query(DailyAttendance)
            .filter(
                DailyAttendance.day >= date_from,
                DailyAttendance.day <= date_to
            )
            .all()
        )
        emp_records_map: dict[str, dict[date, DailyAttendance]] = {}
        for r in daily_records:
            emp_records_map.setdefault(r.emp_no, {})[r.day] = r

        wb = Workbook()

        # 1-varaq: Umumiy
        ws_summary = wb.active
        ws_summary.title = "Umumiy"
        ws_summary.append([
            "ID", "Xodim", "Bo'lim", "Ish boshlanishi",
            "Kelgan kunlar", "Kelmagan kunlar", "Kechikkan kunlar",
            "Jami kechikish (daq)", "O'rtacha kechikish (daq)"
        ])

        # 2-varaq: Kunlik
        ws_daily = wb.create_sheet("Kunlik")
        ws_daily.append([
            "ID", "Xodim", "Sana", "Hafta kuni", "Keldi", "Ketdi", "Ish vaqti", "Kechikish (daq)", "Holat"
        ])

        # 3-varaq: Tabel (Xodim x Kun matritsasi)
        ws_tabel = wb.create_sheet("Tabel")
        date_list = []
        d = date_from
        while d <= date_to:
            date_list.append(d)
            d += timedelta(days=1)

        tabel_header = ["ID", "Xodim", "Bo'lim"] + [f"{cur_d:%d.%m}" for cur_d in date_list] + ["Jami keldi", "Jami kech (daq)"]
        ws_tabel.append(tabel_header)

        for emp in employees:
            records = emp_records_map.get(emp.emp_no, {})
            came_days = 0
            absent_days = 0
            late_days = 0
            total_late = 0

            tabel_row = [emp.emp_no, emp.name, emp.department or ""]

            for cur_d in date_list:
                r = records.get(cur_d)
                wd_name = WEEKDAYS[cur_d.weekday()]

                if r and (r.first_in or r.status in ("present", "late")):
                    came_days += 1
                    late = r.late_minutes
                    if late > 0:
                        late_days += 1
                        total_late += late
                        tabel_row.append(f"K ({late})")
                    else:
                        tabel_row.append("V")

                    ws_daily.append([
                        emp.emp_no,
                        emp.name,
                        f"{cur_d:%d.%m.%Y}",
                        wd_name,
                        f"{r.first_in:%H:%M}" if r.first_in else "",
                        f"{r.last_out:%H:%M}" if r.last_out else "",
                        worked_str(r.worked_minutes),
                        late if late > 0 else "",
                        "Kech keldi" if late > 0 else "Vaqtida keldi"
                    ])
                else:
                    status_str = r.status if r else "Kelmadi"
                    if status_str == "weekend":
                        tabel_row.append("D")
                    elif status_str == "holiday":
                        tabel_row.append("B")
                    else:
                        absent_days += 1
                        tabel_row.append("X")

            tabel_row.append(came_days)
            tabel_row.append(total_late)
            ws_tabel.append(tabel_row)

            avg_late = round(total_late / late_days) if late_days > 0 else 0
            ws_summary.append([
                emp.emp_no,
                emp.name,
                emp.department or "",
                f"{emp.work_start:%H:%M}",
                came_days,
                absent_days,
                late_days,
                total_late,
                avg_late
            ])

        style_sheet(ws_summary)
        style_sheet(ws_daily)
        style_sheet(ws_tabel)

        period = f"{date_from:%d.%m.%Y}_{date_to:%d.%m.%Y}"
        filename = f"davomat_hisobot_barcha_{period}.xlsx"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf, filename
